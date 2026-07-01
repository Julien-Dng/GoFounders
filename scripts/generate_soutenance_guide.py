from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT_DIR = Path("docs")
DOCX_PATH = OUTPUT_DIR / "soutenance_guide_stage_L3.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
TEXT_MUTED = "5D6B7A"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_inches: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(grid_col)


def set_run(run, *, bold=False, italic=False, size=11, color="000000") -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_run(run, bold=True, size=24, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(subtitle)
    set_run(run, size=11, color=TEXT_MUTED)


def add_h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.add_run(text)


def add_h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.add_run(text)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_check(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"[ ] {text}")


def add_callout(doc: Document, label: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(label)
    set_run(run, bold=True, color=DARK_BLUE)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(body)
    doc.add_paragraph()


def add_label_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.65, 4.85])
    hdr = table.rows[0]
    hdr.cells[0].text = "Element"
    hdr.cells[1].text = "A préparer / à renseigner"
    for cell in hdr.cells:
        set_cell_shading(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, bold=True, color=NAVY)

    for label, detail in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = detail
        for cell in row.cells:
            set_cell_margins(cell)
    doc.add_paragraph()


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        set_cell_shading(cell, LIGHT_BLUE)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, bold=True, color=NAVY)

    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
            set_cell_margins(cells[idx])
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.18


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Guide de préparation - stage L3 informatique")
    set_run(run, size=9, color=TEXT_MUTED)


def build_document() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    add_title(
        doc,
        "Guide de préparation - rapport et soutenance de stage L3 informatique",
        "Document de synthèse à compléter pour préparer le rapport/formulaire et le support oral PDF.",
    )

    add_callout(
        doc,
        "Objectif du document",
        "Regrouper les éléments attendus par les consignes de soutenance : informations du stage, missions réalisées, résultats, technologies, bibliographie, plan de présentation et checklist finale.",
    )

    add_h1(doc, "1. Informations à renseigner")
    add_label_table(
        doc,
        [
            ("Stagiaire", "[Nom, prénom, parcours : ASR / CILS / MIAGE]"),
            ("Entreprise", "[Nom de l'entreprise, secteur, activité principale]"),
            ("Tuteur / tutrice", "[Nom, fonction, rôle dans le suivi du stage]"),
            ("Dates", "Mai 2026 - juillet 2026, à ajuster selon la convention."),
            ("Intitulé du stage", "[Titre officiel figurant sur l'offre ou la convention]"),
            ("Sujet initial", "[Sujet présenté au départ, en 4 à 5 lignes rédigées]"),
        ],
    )

    add_h1(doc, "2. Ce que les consignes demandent pour le rapport")
    add_body(
        doc,
        "Le rapport est un formulaire synthétique destiné aux enseignants évaluateurs. Les réponses doivent être rédigées sous forme de phrases, pas sous forme de mots-clés isolés.",
    )
    add_matrix(
        doc,
        ["Rubrique", "Contenu attendu", "Volume conseillé"],
        [
            ["Contexte", "Expliquer l'existant, les contraintes et le cadre du stage.", "5 à 10 lignes"],
            ["Equipe d'accueil", "Présenter l'organisation, les interactions, les outils de travail et le suivi.", "5 à 10 lignes"],
            ["Missions réalisées", "Décrire les missions en ordre chronologique, même si elles ont évolué.", "20 à 30 lignes"],
            ["Missions principales", "Détailler les travaux les plus représentatifs.", "10 à 20 lignes"],
            ["Résultats obtenus", "Présenter les livrables, corrections, fonctionnalités ou apprentissages concrets.", "10 à 20 lignes"],
            ["Bilan et perspectives", "Faire le bilan professionnel et expliquer l'intérêt pour l'entreprise.", "10 à 15 lignes"],
            ["Technologies", "Lister les outils et expliquer leur rôle dans les missions.", "5 à 10 lignes"],
            ["Bibliographie", "Citer documentation, sites, cours, pages techniques utiles, avec commentaire.", "10 à 20 lignes"],
        ],
        [1.35, 3.75, 1.4],
    )

    add_h1(doc, "3. Missions à valoriser dans ton stage")
    add_matrix(
        doc,
        ["Bloc", "Travail à présenter", "Preuves / captures possibles"],
        [
            [
                "Analyse et UX/UI",
                "Recueil des besoins, conception de l'architecture logicielle, réalisation de maquettes Figma.",
                "Maquettes, schéma d'architecture, liste des besoins, choix UX.",
            ],
            [
                "Développement front-end",
                "Intégration d'interfaces utilisateur en HTML5, CSS3, JavaScript et Angular.",
                "Captures d'écran avant/après, composants, pages, responsive.",
            ],
            [
                "Développement back-end",
                "Création de la base de données, développement d'API et de la logique serveur.",
                "Modèle de données, endpoints, flux client/serveur, règles de sécurité.",
            ],
            [
                "Tests et déploiement",
                "Réalisation de tests unitaires, vérifications fonctionnelles et mise en production.",
                "Résultats de tests, procédure de déploiement, environnement cible.",
            ],
        ],
        [1.55, 3.25, 1.7],
    )

    add_h1(doc, "4. Plan conseillé pour la soutenance orale de 15 minutes")
    add_body(
        doc,
        "La soutenance doit être claire, fluide et cohérente. Le support doit rester sobre, numéroté, au format PDF, sans vidéo ni animation. Prévoir des captures plutôt qu'une démonstration en direct.",
    )
    add_matrix(
        doc,
        ["Diapo", "Contenu", "Temps"],
        [
            ["1", "Titre, nom, entreprise, dates, tuteur.", "0:45"],
            ["2", "Contexte de l'entreprise et besoin initial.", "1:15"],
            ["3", "Sujet de stage, objectifs et périmètre.", "1:30"],
            ["4", "Organisation du travail, méthode, outils, interactions.", "1:15"],
            ["5", "Analyse, UX/UI, architecture et maquettes.", "2:00"],
            ["6", "Développement front-end Angular et intégration des interfaces.", "2:00"],
            ["7", "Back-end, base de données, API et logique serveur.", "2:00"],
            ["8", "Tests, correction, déploiement et qualité.", "1:30"],
            ["9", "Résultats obtenus, limites et difficultés rencontrées.", "1:30"],
            ["10", "Bilan professionnel, compétences acquises, perspectives.", "1:15"],
        ],
        [0.65, 4.85, 1.0],
    )

    add_h1(doc, "5. Checklist des éléments à préparer")
    add_h2(doc, "Pour le rapport / formulaire")
    for item in [
        "Rédiger le sujet initial en 5 lignes maximum.",
        "Rédiger le contexte du stage avec l'existant et les contraintes.",
        "Décrire l'équipe d'accueil et les interactions professionnelles.",
        "Lister les missions réalisées dans l'ordre chronologique.",
        "Détailler 2 ou 3 missions principales avec résultats concrets.",
        "Préparer une liste de technologies avec leur rôle précis.",
        "Préparer une bibliographie commentée : Angular, TypeScript, Firebase, documentation API, outils utilisés, cours, etc.",
    ]:
        add_check(doc, item)

    add_h2(doc, "Pour le support de soutenance")
    for item in [
        "Créer un support PDF uniquement, sans animation ni vidéo.",
        "Numéroter toutes les diapositives.",
        "Limiter le texte : mots-clés, captures, schémas, tableaux courts.",
        "Prévoir au moins une capture ou un schéma pour chaque mission importante.",
        "Ne pas faire de démonstration en direct : utiliser des captures commentées.",
        "Répéter la présentation plusieurs fois pour respecter les 15 minutes.",
        "Préparer 5 à 8 questions possibles du jury avec réponses courtes.",
    ]:
        add_check(doc, item)

    add_h1(doc, "6. Formulations courtes réutilisables")
    add_label_table(
        doc,
        [
            (
                "Mission générale",
                "Participation au développement, à l'amélioration et à la mise en qualité d'applications web dans un contexte professionnel.",
            ),
            (
                "Front-end",
                "Intégration d'interfaces utilisateur responsives avec Angular, HTML5, CSS3 et TypeScript.",
            ),
            (
                "Back-end / données",
                "Contribution à la structuration des données, à la logique serveur et aux échanges entre le client et le serveur.",
            ),
            (
                "Qualité",
                "Réalisation de tests, corrections fonctionnelles, vérifications de routage et préparation au déploiement.",
            ),
            (
                "Bilan",
                "Ce stage m'a permis de renforcer mes compétences techniques tout en découvrant les contraintes d'un projet réel : besoins utilisateurs, priorisation, qualité, documentation et livraison.",
            ),
        ],
    )

    add_h1(doc, "7. Questions possibles du jury")
    for question in [
        "Pourquoi avoir choisi cette architecture logicielle ?",
        "Quelles difficultés techniques as-tu rencontrées et comment les as-tu résolues ?",
        "Comment as-tu vérifié que les fonctionnalités dévelopées fonctionnaient correctement ?",
        "Quels choix UX/UI ont été faits pour améliorer l'expérience utilisateur ?",
        "Quelles limites restent à corriger ou améliorer après ton stage ?",
        "Quelles compétences as-tu le plus développées pendant cette période ?",
    ]:
        add_bullet(doc, question)

    doc.add_section(WD_SECTION_START.CONTINUOUS)
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
    print(DOCX_PATH.resolve())
